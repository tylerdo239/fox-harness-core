"""Export one data source's profile as a Word document (Vietnamese) for a data engineer to fill in.

Everything the profile holds about the source goes into editable tables: its tables, every column,
the relationships, the metrics and business terms built on it. Known values are filled in; empty
required cells are highlighted; names that identify things (tables, columns) are greyed out so they
are not changed. Blank rows are left for new relationships, metrics and terms. Nothing is read back:
the filled document is used by a person to configure the profile.
"""

import io
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import _Cell

from src.crud_mongo import entity as entity_crud
from src.crud_mongo import relationship as relationship_crud
from src.data_profile import glossary as glossary_service
from src.data_profile import metrics as metric_service
from src.data_profile import service
from src.data_profile.models import NO_VALUE_OPS, TypedFilter
from src.database.mongodb import AttrDatabase

GREY = "EDEDED"       # names: do not change
YELLOW = "FFF4C2"     # required and still empty
HEAD = "DCE6F2"       # header row
BLANK_ROWS = 5        # empty rows for new relationships, metrics, terms
VN = ZoneInfo("Asia/Ho_Chi_Minh")

TABLE_KIND = {"fact": "fact – sự kiện/giao dịch", "dim": "dim – danh mục", "snapshot": "snapshot – ảnh chụp theo kỳ",
              "scd2": "scd2 – danh mục có lịch sử"}
ROLE = {"key": "key – khóa", "dimension": "dimension – chiều phân tích", "measure": "measure – số đo"}
SEMANTIC = {"currency": "tiền", "date": "ngày", "datetime": "ngày giờ", "category": "phân loại", "id": "mã định danh",
            "percent": "phần trăm", "count": "số lượng", "number": "số (không cộng dồn)", "text": "văn bản",
            "pii": "thông tin cá nhân", "boolean": "đúng/sai"}
AGG = {"count": "Đếm", "count_distinct": "Đếm không trùng", "sum": "Tổng", "avg": "Trung bình", "min": "Nhỏ nhất",
       "max": "Lớn nhất"}
TERM_KIND = {"segment": "segment – nhóm dòng", "metric": "metric – tên gọi khác của chỉ số",
             "definition": "definition – định nghĩa"}
OPS = {"=": "=", "!=": "≠", ">": ">", ">=": "≥", "<": "<", "<=": "≤", "in": "thuộc", "not_in": "không thuộc",
       "contains": "chứa", "like": "giống"}


# ── docx helpers ──

def _shade(cell: _Cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()  # python-docx has no API for cell shading
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _fix_widths(t: Any, widths: list[float]) -> None:
    """Column widths that Word and LibreOffice both keep: fixed layout, the grid and every cell."""
    t.autofit = False
    tbl_pr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for col, w in zip(t.columns, widths, strict=True):
        col.width = Cm(w)
    for row in t.rows:
        for cell, w in zip(row.cells, widths, strict=True):
            cell.width = Cm(w)


def _write(cell: _Cell, text: str, bold: bool = False, size: float = 8.5) -> None:
    cell.text = ""
    lines = text.split("\n") if text else [""]
    p = cell.paragraphs[0]
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.size = Pt(size)
        run.bold = bold
        if i < len(lines) - 1:
            run.add_break()


Col = tuple[str, float, str]   # (header, width cm, kind: "name" grey | "req" yellow when empty | "")


def _table(doc: Any, cols: list[Col], rows: list[list[str]], blank: int = 0) -> None:
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    head = t.rows[0]
    _repeat_header(head)
    for cell, (title, width, kind) in zip(head.cells, cols, strict=True):
        _write(cell, title + (" *" if kind == "req" else ""), bold=True)
        _shade(cell, HEAD)
        cell.width = Cm(width)
    for n, values in enumerate([*rows, *([[""] * len(cols)] * blank)]):
        cells = t.add_row().cells
        for cell, value, (_, width, kind) in zip(cells, values, cols, strict=True):
            _write(cell, value, bold=kind == "name")
            cell.width = Cm(width)
            if kind == "name" and value:
                _shade(cell, GREY)
            elif kind == "req" and not value.strip() and n < len(rows):   # blank rows for new items are optional
                _shade(cell, YELLOW)
    _fix_widths(t, [w for _, w, _ in cols])
    doc.add_paragraph()


def _kv_table(doc: Any, items: list[tuple[str, str, bool]]) -> None:
    """A two-column 'field | value' table; required empty values are highlighted."""
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for label, value, required in items:
        cells = t.add_row().cells
        _write(cells[0], label + (" *" if required else ""), bold=True)
        _shade(cells[0], HEAD)
        _write(cells[1], value)
        cells[0].width, cells[1].width = Cm(6), Cm(19.5)
        if required and not value.strip():
            _shade(cells[1], YELLOW)
    _fix_widths(t, [6, 19.5])
    doc.add_paragraph()


def _heading(doc: Any, text: str, level: int) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)


# ── text of profile values ──

def _filter_text(f: TypedFilter, names: dict[str, str]) -> str:
    col = names.get(f.column_id, "?")
    if f.op in NO_VALUE_OPS:
        return f"{col} {'trống' if f.op == 'is_null' else 'không trống'}"
    return f"{col} {OPS.get(f.op.value, f.op.value)} {', '.join(map(str, f.values))}"


def _filters(filters: list[TypedFilter], names: dict[str, str]) -> str:
    return "\n".join(_filter_text(f, names) for f in filters)


def _values(profile: Any) -> str:
    items = profile.value_catalog
    if not items:
        return ""
    lines = [f"{v.value} = {v.label}" if v.label and v.label != v.value else v.value for v in items]
    return "\n".join(lines)


def _list(xs: list[str] | None) -> str:
    return ", ".join(x for x in (xs or []) if x)


# ── the document ──

def build_docx(db: AttrDatabase, source: dict[str, Any]) -> bytes:
    tables = sorted(entity_crud.list_by_data_source(db, source["_id"]), key=lambda e: e["physical_path"])
    ids = {e["_id"] for e in tables}
    columns = {e["_id"]: service.entity_columns(db, e["_id"]) for e in tables}
    names = {c["_id"]: f"{e['physical_name']}.{c['physical_name']}" for e in tables for c in columns[e["_id"]]}
    short = {c["_id"]: c["physical_name"] for cs in columns.values() for c in cs}

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.5))
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    title = doc.add_heading(f"HỒ SƠ DỮ LIỆU – {source.get('name', '')}", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    doc.add_paragraph(f"Nguồn dữ liệu: {source.get('name', '')} ({source.get('source_type', '')}) · "
                      f"{len(tables)} bảng · Xuất ngày {datetime.now(VN):%d/%m/%Y}")

    _heading(doc, "Hướng dẫn điền", 1)
    for line in (
        "Tài liệu này mô tả các bảng của nguồn dữ liệu để trợ lý AI hiểu và truy vấn đúng. Nhờ anh/chị bổ sung "
        "và sửa những chỗ chưa đúng, rồi gửi lại.",
        "Ô nền xám (tên bảng, tên cột): giữ nguyên, không sửa.",
        "Ô có dấu * là bắt buộc. Ô nền vàng là ô bắt buộc còn trống. Dòng trống ở cuối bảng chỉ điền khi cần thêm mới.",
        "Nội dung đã có là bản nháp: sửa trực tiếp nếu chưa đúng.",
        "Với các ô chọn (loại bảng, vai trò cột…), dùng một trong các giá trị ở bảng giải thích bên dưới.",
        "Danh sách giá trị: mỗi dòng một giá trị, dạng mã = ý nghĩa (ví dụ DONE = Hoàn tất).",
        "Chưa chắc hoặc có câu hỏi: ghi vào cột “Ghi chú / câu hỏi”.",
        "Có thể thêm dòng mới vào các bảng Quan hệ, Chỉ số, Thuật ngữ (đã có sẵn vài dòng trống).",
    ):
        doc.add_paragraph(line, style="List Bullet")

    _heading(doc, "Bảng giải thích lựa chọn", 2)
    _table(doc, [("Mục", 4.5, ""), ("Giá trị và ý nghĩa", 22.2, "")], [
        ["Loại bảng", "fact – mỗi dòng là một sự kiện/giao dịch (đơn hàng, cuộc gọi…)\n"
                      "dim – danh mục, mỗi dòng là một đối tượng (khách hàng, chi nhánh…)\n"
                      "snapshot – trạng thái chụp lại theo ngày/tháng (tồn kho cuối ngày…)\n"
                      "scd2 – danh mục có lưu lịch sử thay đổi (cột hiệu lực từ/đến)"],
        ["Vai trò cột", "key – khóa, dùng để nối bảng / xác định dòng\ndimension – dùng để lọc, nhóm (loại, trạng "
                        "thái, ngày…)\nmeasure – số đo để cộng/đếm/trung bình (số tiền, số lượng…)"],
        ["Loại nghiệp vụ", ", ".join(f"{k} ({v})" for k, v in SEMANTIC.items())],
        ["Ẩn khỏi AI", "Có – cột không dùng để trả lời câu hỏi (cột kỹ thuật, JSON…)\nKhông – cột được dùng"],
        ["Kiểu quan hệ", "1-N – một dòng bảng A ứng với nhiều dòng bảng B\n1-1 – một dòng ứng với đúng một dòng\n"
                         "N-N – nhiều dòng ứng với nhiều dòng"],
        ["Cách tính chỉ số", "Đếm (số dòng), Đếm không trùng, Tổng, Trung bình, Nhỏ nhất, Lớn nhất – của một cột; "
                             "hoặc tỉ lệ = chỉ số A / chỉ số B"],
        ["Loại thuật ngữ", "segment – một nhóm dòng theo điều kiện (khách VIP = hạng = Gold)\n"
                           "metric – tên gọi khác của một chỉ số\ndefinition – định nghĩa, giải thích chung"],
    ])

    # 1. tables
    _heading(doc, "1. Danh sách bảng", 1)
    rows = []
    for i, e in enumerate(tables, 1):
        p = service.entity_profile(e)
        rows.append([
            str(i), e["physical_path"], e.get("display_name") if e.get("display_name") != e["physical_name"] else "",
            e.get("description") or "", e.get("grain_description") or "",
            TABLE_KIND.get(p.table_kind or "", p.table_kind or ""),
            _list([short.get(c, "") for c in p.grain_key_column_ids]), short.get(p.label_column_id or "", ""),
            short.get(p.time_column_id or "", ""), "" if e.get("is_exposed", True) else "Bảng đang ẩn khỏi AI",
        ])
    _table(doc, [("STT", 1.1, ""), ("Bảng", 4.6, "name"), ("Tên hiển thị", 2.8, "req"), ("Mô tả", 4.8, "req"),
                 ("Một dòng là", 3.2, "req"), ("Loại bảng", 2.2, "req"), ("Cột xác định một dòng", 2.4, "req"),
                 ("Cột tên (thay cho mã)", 2, ""), ("Cột thời gian chính", 2, ""), ("Ghi chú / câu hỏi", 1.6, "")],
           rows)

    # 2. each table
    _heading(doc, "2. Chi tiết từng bảng", 1)
    for i, e in enumerate(tables, 1):
        p = service.entity_profile(e)
        _heading(doc, f"2.{i}. {e['physical_name']}", 2)
        doc.add_paragraph(f"Đường dẫn: {e['physical_path']} · {len(columns[e['_id']])} cột")
        _kv_table(doc, [
            ("Múi giờ lưu trữ thời gian", p.storage_tz or "", False),
            ("Múi giờ tính ngày/tháng nghiệp vụ", p.business_tz or "", False),
            ("Dữ liệu có từ ngày", p.coverage_start or "", False),
            ("Dữ liệu đến ngày (trống = đang cập nhật)", p.coverage_end or "", False),
            ("Khoảng thời gian thiếu dữ liệu", p.coverage_gaps or "", False),
            ("Điều kiện luôn áp dụng (vd bỏ dòng đã xóa)", _filters(p.default_filters, names), False),
            ("Lưu ý khi dùng bảng", "\n".join(p.caveats), False),
            ("Ghi chú", p.notes or "", False),
        ])
        rows = []
        for c in columns[e["_id"]]:
            cp = service.column_profile(c)
            rows.append([
                c["physical_name"], c.get("data_type") or "",
                c.get("display_name") if c.get("display_name") != c["physical_name"] else "",
                c.get("description") or "", ROLE.get(c.get("role") or "", c.get("role") or ""),
                SEMANTIC.get(c.get("semantic_type") or "", c.get("semantic_type") or ""), cp.unit or "",
                _values(cp), "Không" if c.get("is_exposed", True) else "Có", cp.notes or "",
            ])
        _table(doc, [("Cột", 3.6, "name"), ("Kiểu", 2.2, "name"), ("Tên hiển thị", 2.8, "req"), ("Mô tả", 4.4, "req"),
                     ("Vai trò", 2.2, "req"), ("Loại nghiệp vụ", 2, ""), ("Đơn vị", 1.4, ""),
                     ("Giá trị có thể có (mã = ý nghĩa)", 4.4, ""), ("Ẩn khỏi AI", 1.5, ""),
                     ("Ghi chú / câu hỏi", 2.2, "")], rows)

    # 3. relationships
    _heading(doc, "3. Quan hệ giữa các bảng", 1)
    doc.add_paragraph("Cách các bảng nối với nhau. Thêm quan hệ còn thiếu vào các dòng trống.")
    rels = relationship_crud.list_touching_entity_ids(db, list(ids))
    pairs = relationship_crud.list_column_pairs_by_relationship_ids(db, [r["_id"] for r in rels])
    others = {x["_id"]: x for x in entity_crud.list_by_ids(db, list({r["from_entity_id"] for r in rels}
                                                                    | {r["to_entity_id"] for r in rels}))}
    rows = []
    for r in rels:
        rp = [x for x in pairs if x["relationship_id"] == r["_id"]]
        a, b = others.get(r["from_entity_id"]), others.get(r["to_entity_id"])
        rows.append([
            a["physical_name"] if a else "?", "\n".join(short.get(x["from_column_id"], "?") for x in rp),
            b["physical_name"] if b else "?", "\n".join(short.get(x["to_column_id"], "?") for x in rp),
            (r.get("cardinality") or "").replace(":", "-"), service.relationship_profile(r).notes or "",
        ])
    _table(doc, [("Bảng A", 4.5, "req"), ("Cột của A", 4, "req"), ("Bảng B", 4.5, "req"), ("Cột của B", 4, "req"),
                 ("Kiểu quan hệ", 2.4, "req"), ("Ghi chú / câu hỏi", 7.3, "")], rows, blank=BLANK_ROWS)

    # 4. metrics
    _heading(doc, "4. Chỉ số", 1)
    doc.add_paragraph("Các con số người dùng hay hỏi (doanh thu, số đơn…) và cách tính. Thêm chỉ số còn thiếu.")
    metric_items = metric_service.list_items(db)
    by_id = {m.id: m for m in metric_items}
    mine = {m.id for m in metric_items if m.table and m.table.entity_id in ids}
    mine |= {m.id for m in metric_items if m.kind == "ratio"
             and {m.numerator_metric_id, m.denominator_metric_id} & mine}
    rows = []
    for m in (by_id[k] for k in by_id if k in mine):
        if m.kind == "ratio":
            calc = f"{m.numerator_name or '?'} / {m.denominator_name or '?'}" + (
                f" × {m.ratio_scale:g}" if m.ratio_scale != 1 else "")
            table = ""
        else:
            what = names.get(m.column_id or "", "") if m.column_id else "số dòng"
            calc = f"{AGG.get(m.aggregation or '', m.aggregation or '')} {what}".strip()
            table = m.table.physical_path if m.table else ""
        rows.append([m.name, m.display_name, m.description or "", table, calc, _filters(m.filters, names), m.unit or "",
                     "\n".join(m.example_questions),
                     ("[đang tắt] " if m.disabled else "") + (m.notes or "")])
    _table(doc, [("Mã chỉ số", 3.4, "req"), ("Tên hiển thị", 2.8, "req"), ("Mô tả", 4, "req"), ("Bảng", 3.2, "req"),
                 ("Cách tính", 3.2, "req"), ("Điều kiện", 3, ""), ("Đơn vị", 1.6, ""), ("Câu hỏi ví dụ", 3, ""),
                 ("Ghi chú / câu hỏi", 2.5, "")], rows, blank=BLANK_ROWS)

    # 5. glossary
    _heading(doc, "5. Thuật ngữ nghiệp vụ", 1)
    doc.add_paragraph("Từ ngữ riêng của doanh nghiệp và ý nghĩa của chúng. Thêm thuật ngữ còn thiếu.")
    rows = []
    for t in glossary_service.list_items(db):
        linked = ({t.table.entity_id} if t.table else set()) | {r.entity_id for r in t.related_tables}
        if linked and not linked & ids and t.metric_id not in mine:
            continue    # a term of another source
        if t.kind == "segment":
            applies = f"{t.table.physical_path if t.table else '?'}: {_filters(t.filters, names)}"
        elif t.kind == "metric":
            applies = f"chỉ số {t.metric_name or '?'}"
        else:
            applies = _list([r.physical_path for r in t.related_tables])
        rows.append([t.term, _list(t.synonyms), TERM_KIND.get(t.kind, t.kind), t.definition or "", applies,
                     "\n".join(t.example_questions), ("[đang tắt] " if t.disabled else "") + (t.notes or "")])
    _table(doc, [("Thuật ngữ", 3.2, "req"), ("Cách gọi khác", 3.4, ""), ("Loại", 2.8, "req"), ("Định nghĩa", 6, "req"),
                 ("Áp dụng cho (bảng + điều kiện / chỉ số)", 5, ""), ("Câu hỏi ví dụ", 3.3, ""),
                 ("Ghi chú / câu hỏi", 3, "")], rows, blank=BLANK_ROWS)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
